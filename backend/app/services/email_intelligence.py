import io
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

import fitz
from docx import Document

from ..schemas import (
    DeadlineModel,
    EmailFieldConfidence,
    InterviewEmailAnalysis,
    InterviewerModel,
    PreparationChecklistModel,
    SuggestedPreparationModel,
    TimelineStepModel,
)


PLATFORM_URL_PATTERNS = [
    ("Google Meet", re.compile(r"https?://meet\.google\.com/[a-z]{3}-[a-z]{4}-[a-z]{3}\b", re.I)),
    ("Zoom", re.compile(r"https?://(?:[a-zA-Z0-9-]+\.)?zoom\.us/(?:j/|my/|wc/)[0-9a-zA-Z?=&_#-]+", re.I)),
    ("Microsoft Teams", re.compile(r"https?://teams\.microsoft\.com/(?:l/meetup-join/[^\s<>\"']+)", re.I)),
    ("Webex", re.compile(r"https?://[a-zA-Z0-9-]+\.webex\.com/[^\s<>\"']+", re.I)),
]

INTERVIEW_TYPES = [
    ("Technical Interview", re.compile(r"\btechnical\s+(?:interview|round|discussion)\b", re.I)),
    ("Coding Interview", re.compile(r"\b(?:coding|live\s+coding|programming|algorithm)\s+(?:interview|round|assessment)\b", re.I)),
    ("HR Interview", re.compile(r"\b(?:hr|human\s+resources|people)\s+(?:interview|round|discussion)\b", re.I)),
    ("Behavioral Interview", re.compile(r"\b(?:behavioral|culture\s+fit|values)\s+(?:interview|round|discussion)\b", re.I)),
    ("Managerial Interview", re.compile(r"\b(?:managerial|hiring\s+manager|engineering\s+manager)\s+(?:interview|round)\b", re.I)),
    ("Screening", re.compile(r"\b(?:screening|initial\s+screening|introductory\s+call|screen)\b", re.I)),
    ("Assessment", re.compile(r"\b(?:assessment|online\s+test|take-home|hackerrank|leetcode)\b", re.I)),
    ("Panel Interview", re.compile(r"\b(?:panel\s+interview|panel\s+round)\b", re.I)),
    ("Video Interview", re.compile(r"\b(?:video\s+interview|virtual\s+interview)\b", re.I)),
    ("Phone Interview", re.compile(r"\b(?:phone\s+interview|telephonic\s+interview|phone\s+screen)\b", re.I)),
    ("On-site Interview", re.compile(r"\b(?:on-site|onsite|in-person|office\s+interview)\b", re.I)),
]

MONTH_NAMES = (
    "January|February|March|April|May|June|July|August|September|October|November|December|"
    "Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
)


def extract_email_text_from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".eml", ".email"):
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            return data.decode("latin1", errors="ignore")
    if suffix == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as doc:
            pages = [page.get_text("text", sort=True) for page in doc]
            return "\n".join(pages).strip()
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = [cell.text for t in doc.tables for row in t.rows for cell in row.cells if cell.text.strip()]
        return "\n".join(paragraphs + tables).strip()
    raise ValueError("Only TXT, PDF, and DOCX files are supported for email analysis.")


def extract_meeting_link(text: str) -> tuple[str | None, str | None]:
    for platform_name, pattern in PLATFORM_URL_PATTERNS:
        match = pattern.search(text)
        if match:
            return match.group(0), platform_name

    generic_url = re.search(r"https?://[^\s<>\"']+(?:join|meet|interview|call)[^\s<>\"']*", text, re.I)
    if generic_url:
        return generic_url.group(0), None
    return None, None


def extract_headers(text: str) -> dict[str, str | None]:
    headers: dict[str, str | None] = {"subject": None, "sender": None, "recipient": None, "date": None}
    for line in text.splitlines()[:20]:
        line = line.strip()
        if not line:
            continue
        subj_m = re.match(r"^(?:Subject|Sub):\s*(.+)$", line, re.I)
        if subj_m and not headers["subject"]:
            headers["subject"] = subj_m.group(1).strip()
        from_m = re.match(r"^(?:From|Sender):\s*(.+)$", line, re.I)
        if from_m and not headers["sender"]:
            headers["sender"] = from_m.group(1).strip()
        to_m = re.match(r"^(?:To|Recipient):\s*(.+)$", line, re.I)
        if to_m and not headers["recipient"]:
            headers["recipient"] = to_m.group(1).strip()
        date_m = re.match(r"^(?:Date|Sent):\s*(.+)$", line, re.I)
        if date_m and not headers["date"]:
            headers["date"] = date_m.group(1).strip()
    return headers


def extract_company(text: str, sender: str | None = None) -> EmailFieldConfidence:
    explicit_m = re.search(r"\bCompany\s*:\s*([A-Za-z0-9&.\s]{2,40})", text, re.I)
    if explicit_m:
        val = explicit_m.group(1).strip()
        val = re.split(r"[\n\r,;]", val)[0].strip()
        if val:
            return EmailFieldConfidence(value=val, confidence="clearly_found", source_text=explicit_m.group(0).strip())

    # Check "@ CompanyName" in subject or headers or body
    at_symbol_m = re.search(r"@\s*([A-Z][A-Za-z0-9&]+(?:[ \t]+[A-Z][A-Za-z0-9&]+)*)", text)
    if at_symbol_m:
        val = at_symbol_m.group(1).strip()
        if val.lower() not in ("gmail", "yahoo", "outlook", "google meet"):
            return EmailFieldConfidence(value=val, confidence="clearly_found", source_text=at_symbol_m.group(0).strip())

    # Check "role/interview at/with Company" without crossing lines
    phrase_m = re.search(r"\b(?:interview\s+(?:at|with)|position\s+at|role\s+at|screening\s+(?:at|with)|call\s+(?:at|with)|team\s+at|welcome\s+to)\s+([A-Z][A-Za-z0-9&]+(?:[ \t]+[A-Z][A-Za-z0-9&]+)*)", text, re.I)
    if phrase_m:
        cand = phrase_m.group(1).strip()
        if cand.lower() not in ("google meet", "microsoft teams", "zoom", "our office", "the team", "phone"):
            return EmailFieldConfidence(value=cand, confidence="clearly_found", source_text=phrase_m.group(0).strip())

    if sender:
        domain_m = re.search(r"@([a-zA-Z0-9-]+)\.[a-zA-Z]{2,}", sender)
        if domain_m:
            domain = domain_m.group(1).lower()
            if domain not in ("gmail", "yahoo", "outlook", "hotmail", "icloud", "proton", "protonmail"):
                exact_case_m = re.search(rf"\b({re.escape(domain)})\b", text, re.I)
                formatted = exact_case_m.group(1) if exact_case_m else "".join(word.capitalize() for word in domain.split("-"))
                return EmailFieldConfidence(value=formatted, confidence="inferred", source_text=f"Derived from sender domain: {domain}")

    return EmailFieldConfidence(value="Unknown Company", confidence="not_found")


def extract_role(text: str, candidate_role: str | None = None) -> EmailFieldConfidence:
    explicit_m = re.search(r"\bRole\s*:\s*([A-Za-z0-9&/\s-]{3,45})", text, re.I)
    if explicit_m:
        val = explicit_m.group(1).strip()
        val = re.split(r"[\n\r,;]", val)[0].strip()
        if val:
            return EmailFieldConfidence(value=val, confidence="clearly_found", source_text=explicit_m.group(0).strip())

    for role_name in (
        "Machine Learning Intern", "Backend Developer", "Frontend Developer",
        "Full Stack Developer", "Software Engineer", "Data Scientist", "Data Analyst",
        "DevOps Engineer", "Cloud Architect", "AI Research Intern", "Systems Engineer",
    ):
        if re.search(rf"\b{re.escape(role_name)}\b", text, re.I):
            return EmailFieldConfidence(value=role_name, confidence="clearly_found", source_text=f"Matched role: {role_name}")

    pattern = re.search(r"\b(?:for\s+the|as\s+a[n]?|position\s+of|role\s+of)\s+([A-Za-z0-9&/\s-]{3,35}?(?:Intern|Developer|Engineer|Architect|Scientist|Analyst|Lead|Specialist|Manager))\b", text, re.I)
    if pattern:
        return EmailFieldConfidence(value=pattern.group(1).strip(), confidence="inferred", source_text=pattern.group(0).strip())

    if candidate_role:
        return EmailFieldConfidence(value=candidate_role, confidence="inferred", source_text="Inferred from candidate target role")

    return EmailFieldConfidence(value="Candidate Role", confidence="not_found")


def extract_interview_type(text: str) -> EmailFieldConfidence:
    for type_name, pattern in INTERVIEW_TYPES:
        match = pattern.search(text)
        if match:
            return EmailFieldConfidence(value=type_name, confidence="clearly_found", source_text=match.group(0).strip())

    if "meet" in text.lower() or "zoom" in text.lower() or "teams" in text.lower():
        return EmailFieldConfidence(value="Video Interview", confidence="inferred", source_text="Inferred from video platform presence")

    return EmailFieldConfidence(value="Technical Interview", confidence="inferred", source_text="Default technical assessment route")


def extract_date_and_time(text: str, sent_date: str | None = None) -> tuple[EmailFieldConfidence, EmailFieldConfidence, EmailFieldConfidence]:
    date_val = None
    date_conf = "not_found"
    date_src = None

    # Separate body from initial header block so sent_date doesn't get confused for interview date
    lines = text.splitlines()
    body_start_idx = 0
    in_header = True
    for i, line in enumerate(lines[:15]):
        if in_header and not line.strip():
            body_start_idx = i + 1
            break
        if not re.match(r"^(?:Subject|From|To|Date|Sent|Cc|Bcc):\s*", line, re.I):
            in_header = False
            body_start_idx = i
            break
    body_text = "\n".join(lines[body_start_idx:]) if body_start_idx > 0 else text

    # Search for explicit Date in body first
    explicit_date = re.search(r"\bDate\s*:\s*([^\n\r]+)", body_text, re.I)
    if explicit_date:
        candidate = explicit_date.group(1).strip()
        date_match = re.search(rf"\b(?:{MONTH_NAMES})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s*(?:20\d{{2}})?\b", candidate, re.I)
        if date_match:
            date_val = date_match.group(0).strip()
            date_conf = "clearly_found"
            date_src = explicit_date.group(0).strip()
        elif candidate:
            date_val = candidate[:30].strip()
            date_conf = "clearly_found"
            date_src = explicit_date.group(0).strip()

    if not date_val:
        date_pattern = re.search(rf"\b(?:{MONTH_NAMES})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s*(?:20\d{{2}})?\b", body_text, re.I)
        if date_pattern:
            matched_date = date_pattern.group(0).strip()
            # Ensure it is not identical to the email sent date header if it's the only one
            if not sent_date or matched_date.lower() not in sent_date.lower():
                date_val = matched_date
                date_conf = "clearly_found"
                date_src = date_pattern.group(0).strip()

    time_val = None
    time_conf = "not_found"
    time_src = None

    time_pattern = re.search(r"\b(\d{1,2}(?::\d{2})?\s*(?:AM|PM|am|pm))\b", body_text)
    if time_pattern:
        time_val = time_pattern.group(1).upper().strip()
        time_conf = "clearly_found"
        time_src = time_pattern.group(0).strip()

    tz_val = "IST"
    tz_conf = "inferred"
    tz_src = "Inferred default timezone"

    tz_pattern = re.search(r"\b(IST|PST|PDT|EST|EDT|CST|CDT|UTC|GMT|CET|BST)\b", body_text)
    if tz_pattern:
        tz_val = tz_pattern.group(1).upper()
        tz_conf = "clearly_found"
        tz_src = tz_pattern.group(0)

    date_res = EmailFieldConfidence(value=date_val or "Not specified", confidence=date_conf, source_text=date_src)
    time_res = EmailFieldConfidence(value=time_val or "Not specified", confidence=time_conf, source_text=time_src)
    tz_res = EmailFieldConfidence(value=tz_val, confidence=tz_conf, source_text=tz_src)
    return date_res, time_res, tz_res


def extract_platform(text: str, detected_url_platform: str | None) -> EmailFieldConfidence:
    if detected_url_platform:
        return EmailFieldConfidence(value=detected_url_platform, confidence="clearly_found", source_text=f"Detected meeting link for {detected_url_platform}")

    for p in ("Google Meet", "Zoom", "Microsoft Teams", "Webex"):
        if re.search(rf"\b{re.escape(p)}\b", text, re.I):
            return EmailFieldConfidence(value=p, confidence="clearly_found", source_text=f"Mentioned in email text: {p}")

    if re.search(r"\b(?:phone\s+call|telephonic|phone)\b", text, re.I):
        return EmailFieldConfidence(value="Phone", confidence="clearly_found", source_text="Telephonic interview specified")

    if re.search(r"\b(?:in-person|on-site|office\s+location|visit\s+our\s+office)\b", text, re.I):
        return EmailFieldConfidence(value="In-person", confidence="clearly_found", source_text="In-person location specified")

    return EmailFieldConfidence(value="Unknown", confidence="not_found")


def extract_duration(text: str) -> str | None:
    dur_m = re.search(r"\b(?:Duration\s*:\s*)?(\d{1,3})(?:\s*-\s*\d{1,3})?[-\s]*(mins?|minutes?|hours?|hrs?)\b", text, re.I)
    if dur_m:
        num = dur_m.group(1)
        unit = dur_m.group(2).lower()
        if "hour" in unit or "hr" in unit:
            return f"{num} hour" if num == "1" else f"{num} hours"
        return f"{num} minutes"
    return None


def extract_interviewers(text: str) -> list[InterviewerModel]:
    interviewers: list[InterviewerModel] = []
    explicit_m = re.search(r"\b(?:Interviewer[s]?|Speaking\s+with|Hosted\s+by|With)\s*:\s*([^\n\r]+)", text, re.I)
    if explicit_m:
        chunk = explicit_m.group(1).strip()
        parts = chunk.split(",")
        name = parts[0].strip()
        role = parts[1].strip() if len(parts) > 1 else None
        if name and len(name) < 50:
            interviewers.append(InterviewerModel(name=name, role=role))

    if not interviewers:
        dr_m = re.search(r"\b(Dr\.\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)\b", text)
        if dr_m:
            interviewers.append(InterviewerModel(name=dr_m.group(1).strip(), role="Interviewer"))
    return interviewers


def extract_preparation(text: str, role_title: str) -> tuple[list[str], list[PreparationChecklistModel]]:
    reqs: list[str] = []
    checklist: list[PreparationChecklistModel] = []
    chk_idx = 1

    low = text.lower()
    if "resume" in low:
        reqs.append("Have a copy of your resume ready")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Resume copy ready", done=True))
        chk_idx += 1

    if any(k in low for k in ("portfolio", "github", "project", "code sample")):
        reqs.append("Review previous projects and be prepared to explain architecture")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Review project architecture & code", done=False))
        chk_idx += 1

    if any(k in low for k in ("coding", "editor", "live coding", "problem-solving", "hackerrank")):
        reqs.append("Have your code editor and live coding environment set up")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Coding environment & editor ready", done=False))
        chk_idx += 1

    if any(k in low for k in ("id", "identification", "photo id", "student id", "passport", "aadhaar")):
        reqs.append("Bring a valid government or student identification document")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Government / Student ID available", done=False))
        chk_idx += 1

    if any(k in low for k in ("presentation", "slide", "deck")):
        reqs.append("Prepare presentation deck and slides")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Presentation deck prepared", done=False))
        chk_idx += 1

    if any(k in low for k in ("research", "company information", "about us")):
        reqs.append("Review company background and products")
        checklist.append(PreparationChecklistModel(id=f"chk-{chk_idx}", label="Company research completed", done=False))
        chk_idx += 1

    if not reqs:
        reqs = ["Review your resume and prepare to explain your key technical projects"]
        checklist = [PreparationChecklistModel(id="chk-1", label="Resume and projects review", done=True)]

    return reqs, checklist


def build_suggested_preparation(role: str, candidate_skills: list[str] | None) -> SuggestedPreparationModel:
    candidate_skills = candidate_skills or ["Python", "FastAPI", "REST", "SQL", "Git"]
    low = role.lower()

    if "machine learning" in low or "ai" in low or "data science" in low:
        tech = [
            "Python data structures, NumPy, and Pandas fundamentals",
            "Model training, evaluation metrics (precision, recall, F1, latency)",
            "FastAPI / Flask inference pipeline integration",
            "SQL data querying and aggregation",
        ]
        role_prep = [
            "Explain model baseline, features, and trade-offs on your AI project",
            "Walk through latency reduction or model quantization techniques",
            "Discuss handling messy training data and edge-case inputs",
        ]
    elif "frontend" in low or "react" in low:
        tech = [
            "React 19 hooks, component lifecycle, and state optimization",
            "TypeScript interfaces, generics, and strict type safety",
            "CSS layout, responsive design, and accessible ARIA attributes",
            "Client performance, bundle splitting, and network caching",
        ]
        role_prep = [
            "Explain component breakdown and state management decisions",
            "Walk through accessibility considerations and responsive edge cases",
        ]
    else:  # Default to backend / software engineering
        tech = [
            "Python / FastAPI endpoint design and middleware",
            "REST API contracts, error statuses, and payload validation",
            "Database indexing, transactions, and SQL query profiling",
            "Containerization with Docker and deployment pipelines",
        ]
        role_prep = [
            "Explain latency reduction and caching decisions on your backend project",
            "Walk through handling high concurrency and error recovery",
            "Discuss testing strategies (unit, integration, mock test fixtures)",
        ]

    company_prep = [
        "Research the company's recent engineering and product announcements",
        "Understand the target problem domain and user base",
        "Prepare 2-3 technical questions for the interviewer regarding team architecture",
    ]

    matched = [s for s in candidate_skills if any(s.lower() in t.lower() for t in tech) or s.lower() in role.lower()]
    if not matched:
        matched = candidate_skills[:3]

    return SuggestedPreparationModel(
        technical=tech,
        role=role_prep,
        company=company_prep,
        profile_matched_skills=matched,
    )


def extract_deadlines(text: str) -> list[DeadlineModel]:
    deadlines: list[DeadlineModel] = []
    m = re.search(r"\b(?:confirm|confirm availability|rsvp|submit|respond|complete\s+by)\s+(?:by|before|on)?\s*([A-Za-z0-9,.\s]{4,30}?20\d{2}|[A-Za-z]+\s+\d{1,2})\b", text, re.I)
    if m:
        deadlines.append(DeadlineModel(label="Confirm availability / RSVP", date=m.group(1).strip()))
    return deadlines


def build_timeline(date_val: str, headers_date: str | None, deadlines: list[DeadlineModel]) -> list[TimelineStepModel]:
    steps: list[TimelineStepModel] = [
        TimelineStepModel(
            label="Email Received",
            date=headers_date or "Recently delivered",
            status="completed",
            detail="Invitation received and analyzed",
        )
    ]
    for d in deadlines:
        steps.append(TimelineStepModel(
            label=d.label,
            date=d.date,
            status="completed" if "confirm" in d.label.lower() else "current",
            detail="Action required",
        ))

    steps.append(TimelineStepModel(
        label="Interview Session",
        date=date_val if date_val != "Not specified" else "Date pending",
        status="current",
        detail="Live interview session",
    ))
    steps.append(TimelineStepModel(
        label="Follow-up & Outcome",
        date="1–3 days post-interview",
        status="upcoming",
        detail="Feedback and next round notification",
    ))
    return steps


def analyze_email_text(
    text: str,
    file_name: str | None = None,
    candidate_skills: list[str] | None = None,
    candidate_role: str | None = None,
    source: str = "paste",
) -> InterviewEmailAnalysis:
    text_clean = text.strip()
    if not text_clean:
        raise ValueError("Email content cannot be empty.")

    headers = extract_headers(text_clean)
    meeting_link, detected_url_platform = extract_meeting_link(text_clean)

    company_conf = extract_company(text_clean, headers.get("sender"))
    role_conf = extract_role(text_clean, candidate_role)
    type_conf = extract_interview_type(text_clean)
    date_conf, time_conf, tz_conf = extract_date_and_time(text_clean, headers.get("date"))
    platform_conf = extract_platform(text_clean, detected_url_platform)
    duration = extract_duration(text_clean)
    interviewers = extract_interviewers(text_clean)
    prep_reqs, checklist = extract_preparation(text_clean, role_conf.value)
    suggested_prep = build_suggested_preparation(role_conf.value, candidate_skills)
    deadlines = extract_deadlines(text_clean)
    timeline = build_timeline(date_conf.value, headers.get("date"), deadlines)

    # Location
    location = None
    loc_m = re.search(r"\b(?:Location|Address|Venue)\s*:\s*([^\n\r]+)", text_clean, re.I)
    if loc_m:
        location = loc_m.group(1).strip()

    # Determine status
    status: str = "Upcoming"
    if date_conf.confidence == "not_found":
        status = "Needs Review"

    # Human-readable summary
    interviewer_clause = f" with {interviewers[0].name}" if interviewers else ""
    duration_clause = f" The interview is expected to last {duration} and includes a {type_conf.value.lower()}." if duration else f" Includes a {type_conf.value.lower()}."
    time_part = f" at {time_conf.value} {tz_conf.value}" if time_conf.value != "Not specified" else ""
    date_part = f" on {date_conf.value}" if date_conf.value != "Not specified" else ""
    platform_part = f" via {platform_conf.value}" if platform_conf.value != "Unknown" else ""

    summary = (
        f"Your {role_conf.value} interview with {company_conf.value} is scheduled{date_part}{time_part}{platform_part}{interviewer_clause}.{duration_clause}"
    )

    now_iso = datetime.now(timezone.utc).isoformat()
    email_id = f"ie-{uuid.uuid4().hex[:8]}"

    return InterviewEmailAnalysis(
        id=email_id,
        source=source if source in ("upload", "paste") else "paste",
        file_name=file_name,
        original_email=text_clean,
        subject=headers.get("subject"),
        sender=headers.get("sender"),
        recipient=headers.get("recipient"),
        company=company_conf,
        role=role_conf,
        interview_type=type_conf,
        interview_date=date_conf,
        interview_time=time_conf,
        timezone=tz_conf,
        platform=platform_conf,
        meeting_link=meeting_link,
        location=location,
        duration=duration,
        interviewers=interviewers,
        preparation_requirements=prep_reqs,
        preparation_checklist=checklist,
        suggested_preparation=suggested_prep,
        deadlines=deadlines,
        attachments=["None detected"],
        timeline=timeline,
        summary=summary,
        status=status,  # type: ignore
        created_at=now_iso,
        updated_at=now_iso,
    )
