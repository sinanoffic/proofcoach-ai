import io
import re
from pathlib import Path

import fitz
from docx import Document

from ..schemas import ParserAnalysis, ParserCheck


SKILL_VOCABULARY = [
    "Python", "FastAPI", "REST", "SQL", "Docker", "Testing", "System Design",
    "React", "TypeScript", "JavaScript", "Git", "Machine Learning", "Pandas",
    "NumPy", "PostgreSQL", "AWS", "Linux", "CI/CD",
]
SECTION_PATTERNS = {
    "education": r"\b(education|academic background)\b",
    "skills": r"\b(skills|technical skills|technologies)\b",
    "projects": r"\b(projects?|selected work)\b",
    "experience": r"\b(experience|employment|internships?)\b",
    "certifications": r"\b(certifications?|courses?)\b",
    "achievements": r"\b(achievements?|awards?)\b",
}


def extract_text(data: bytes, filename: str) -> tuple[str, dict[str, bool]]:
    suffix = Path(filename).suffix.lower()
    layout = {"two_column": False, "tables": False}
    if suffix == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            text_parts: list[str] = []
            for page in document:
                text_parts.append(page.get_text("text", sort=True))
                blocks = page.get_text("blocks")
                midpoint = page.rect.width / 2
                left = any(block[0] < midpoint and block[2] < midpoint * 1.15 for block in blocks)
                right = any(block[0] > midpoint * .85 for block in blocks)
                layout["two_column"] = layout["two_column"] or (left and right)
            return "\n".join(text_parts).strip(), layout
    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        layout["tables"] = bool(document.tables)
        return "\n".join(paragraphs + table_text).strip(), layout
    raise ValueError("Only PDF and DOCX files are supported.")


def extract_claims(text: str) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    impact = re.compile(
        r"(\b\d+(?:\.\d+)?\s*%|\b\d+[kKmM]?\+?\s+(?:users?|requests?|clients?)|"
        r"\b(?:reduced|increased|improved|accelerated|saved|grew|led|managed|scaled|optimized)\b)",
        re.IGNORECASE,
    )
    return [sentence.strip(" •-\t") for sentence in sentences if impact.search(sentence)][:12]


def _find_skills(text: str) -> list[str]:
    return [skill for skill in SKILL_VOCABULARY if re.search(rf"\b{re.escape(skill)}\b", text, re.I)]


def analyze_resume_text(text: str, filename: str = "resume.pdf", layout: dict[str, bool] | None = None) -> ParserAnalysis:
    layout = layout or {"two_column": False, "tables": False}
    normalized = text.strip()
    found_sections = {name: bool(re.search(pattern, normalized, re.I)) for name, pattern in SECTION_PATTERNS.items()}
    skills = _find_skills(normalized)
    claims = extract_claims(normalized)
    contact_ok = bool(re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", normalized)) and bool(re.search(r"(?:\+?\d[\s-]?){8,14}", normalized))
    dates_ok = bool(re.search(r"\b(?:19|20)\d{2}\b|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b", normalized, re.I))
    extraction_ratio = min(1.0, len(normalized) / 900)

    checks = [
        ParserCheck(label="Text extraction", status="pass" if extraction_ratio >= .75 else "warning", detail=f"{len(normalized):,} readable characters", points=20 if extraction_ratio >= .75 else round(20 * extraction_ratio), maximum=20),
        ParserCheck(label="Contact details", status="pass" if contact_ok else "warning", detail="Email and phone detected" if contact_ok else "Check email/phone formatting", points=10 if contact_ok else 5, maximum=10),
        ParserCheck(label="Recognizable sections", status="pass" if sum(found_sections.values()) >= 3 else "warning", detail=f"{sum(found_sections.values())}/6 core sections detected", points=min(15, sum(found_sections.values()) * 3), maximum=15),
        ParserCheck(label="Dates", status="pass" if dates_ok else "warning", detail="Timeline markers detected" if dates_ok else "No clear year/month markers", points=10 if dates_ok else 4, maximum=10),
        ParserCheck(label="Skills", status="pass" if len(skills) >= 4 else "warning", detail=f"{len(skills)} recognized skills", points=min(15, len(skills) * 3), maximum=15),
        ParserCheck(label="Projects", status="pass" if found_sections["projects"] else "missing", detail="Project section recovered" if found_sections["projects"] else "Project heading not recovered", points=10 if found_sections["projects"] else 0, maximum=10),
        ParserCheck(label="Two-column layout", status="warning" if layout.get("two_column") else "pass", detail="Reading order may be ambiguous" if layout.get("two_column") else "No split-column warning", points=4 if layout.get("two_column") else 10, maximum=10),
        ParserCheck(label="Table content", status="warning" if layout.get("tables") else "pass", detail="Tables may reorder content" if layout.get("tables") else "No table warning", points=4 if layout.get("tables") else 10, maximum=10),
    ]
    score = max(0, min(100, sum(item.points for item in checks)))
    sections = {name: "detected" if found else "not detected" for name, found in found_sections.items()}
    return ParserAnalysis(
        filename=filename,
        text_preview=normalized[:800],
        sections=sections,
        skills=skills,
        technologies=skills,
        measurable_claims=claims,
        checks=checks,
        parser_robustness=score,
    )


def parse_resume(data: bytes, filename: str) -> ParserAnalysis:
    text, layout = extract_text(data, filename)
    if not text:
        raise ValueError("No readable text was found in the document.")
    return analyze_resume_text(text, filename, layout)

